import pandas as pd
import numpy as np
import mysql.connector
import matplotlib.pyplot as plt
import seaborn as sns
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Cm,Pt
from pptx import chart
from pptx.enum.chart import XL_LEGEND_POSITION,XL_LABEL_POSITION
from pptx.chart.data import ChartData
from pptx.enum.shapes import MSO_SHAPE
from pandas import DataFrame as DF
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LABEL_POSITION
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

""" ****************************************************************** CONECTION TO DB- MYSQL *************************************************************************"""
mydb = mysql.connector.connect(
    host="bizly-db-prod-reader-do-user-12143907-0.b.db.ondigitalocean.com",
    database='defaultdb',
    user="qbr-reader",
    password="AVNS__uHZWnkHP0xksBmbFcL",
    port=25060  
    )

mydb.connect(db='bizly_prod')

sql_query_meetings="""
select  
	"meetings"."team_id", 
    year("meetings"."created_at") AS "year",
    month("meetings"."created_at") AS "month",
    count("meetings"."created_at") as Created, 
    count("meetings"."ends_at") as Completed, 
    count("meetings"."cancelled_at") as Cancelled,
    avg(TIMESTAMPDIFF(week,"meetings"."created_at","meetings"."starts_at")) as lead_time
from "meetings"
where meetings.deleted_at is null
group by "meetings"."team_id",month("meetings"."created_at"),year("meetings"."created_at")
"""

mycursor=mydb.cursor()
mycursor.execute(sql_query_meetings)
myresult = mycursor.fetchall()

meetings=pd.DataFrame(myresult,columns=['team_id','year','month','Created','Completed','Cancelled','lead_time'])
meetings['team_id']=meetings['team_id'].astype('string')

sql_query_proposals="""
select  
	"proposals"."team_id", 
    year("proposals"."created_at") AS "year",
    month("proposals"."created_at") AS "month",
    count("proposals"."accepted_at") as Accepted
from "proposals"
group by "proposals"."team_id",month("proposals"."created_at"),year("proposals"."created_at")
"""

mycursor=mydb.cursor()
mycursor.execute(sql_query_proposals)
myresult = mycursor.fetchall()

proposals=pd.DataFrame(myresult,columns=['team_id','year','month','Accepted'])
proposals['team_id']=proposals['team_id'].astype('string')


meetings=pd.merge(meetings,proposals, how='left' ,left_on=['team_id','month','year'], right_on = ['team_id','month','year'])

sql_query_bookings="""
select 
	"booking_inquiries"."team_id", 
    year("booking_inquiries"."submitted_at") AS "year",
    month("booking_inquiries"."submitted_at") AS "month",
    count("booking_inquiries"."submitted_at") as Submitted,
    avg("booking_inquiry_venues"."response_time") as response_time
from ("booking_inquiries" left join "booking_inquiry_venues" on("booking_inquiries"."id" = "booking_inquiry_venues"."id"))
group by "booking_inquiries"."team_id",month("booking_inquiries"."submitted_at"),year("booking_inquiries"."submitted_at")
"""

mycursor=mydb.cursor()
mycursor.execute(sql_query_bookings)
myresult = mycursor.fetchall()

bookings=pd.DataFrame(myresult,columns=['team_id','year','month','Submitted','response_time'])
bookings['team_id']=bookings['team_id'].astype('string')

meetings=pd.merge(meetings,bookings, how='left' ,left_on=['team_id','month','year'], right_on = ['team_id','month','year'])

sql_query_orders="""

select  
	convert("orders"."team_id", char) as team_id,
    year("orders"."contracted_at") AS "year",
    month("orders"."contracted_at") AS "month",
    count("orders"."contracted_at") as Booked
from "orders"
group by "orders"."team_id",month("orders"."contracted_at"),year("orders"."contracted_at")

"""

mycursor=mydb.cursor()
mycursor.execute(sql_query_orders)
myresult = mycursor.fetchall()

orders=pd.DataFrame(myresult,columns=['team_id','year','month','Booked'])
orders['team_id']=orders['team_id'].astype('string')

meetings=pd.merge(meetings,orders, how='left' ,left_on=['team_id','month','year'], right_on = ['team_id','month','year'])
meetings=meetings.fillna(0)

#client_list=['419']
client_list=['419',	'802',	'10153',	'10982',	'10983',	'11789',	'12693',	'12801',	'13429',	'13449',	'14476',	'15224',	'15356',	'15395']


for client in client_list:

    meetings_client=meetings[meetings['team_id']==client]

    meetings_client=meetings_client[meetings_client['year']==2022]

    if client=='419':
        all_meetings=meetings_client
    else:
        all_meetings=all_meetings.append(meetings_client)




        #ws.cell(row,1).val